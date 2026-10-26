import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import Counter

wb = openpyxl.load_workbook(r'c:\Users\hu\.trae-cn\attachments\6a9fb01e4b6a3916ac6bd24c\6e17faf1-c8a8-4792-8b57-4294fdbe8e4c_f5e5b156-86d5-4fa4-ae36-b7ff7fd8ae27_最终筛选_冯国栋可报岗位.xlsx')
ws = wb.active

# Get headers
headers = [cell.value for cell in ws[1]]
print("Headers:", [(i, h) for i, h in enumerate(headers) if h])

# Find subject column
subject_col = None
for i, h in enumerate(headers):
    if h and '学科专业' in str(h):
        subject_col = i
        break

# Find other useful columns
col_map = {}
for i, h in enumerate(headers):
    if h:
        for key in ['岗位编号', '岗位名称', '从事工作', '学科专业', '工作地点', '用人单位', '大单位', '性别', '学历层次', '院校']:
            if key in str(h):
                col_map[key] = i

print(f"Column map: {col_map}")

# Filter
filtered = []
for row in ws.iter_rows(min_row=2, values_only=True):
    subject = str(row[subject_col]) if row[subject_col] else ''
    if '计算机科学与技术' in subject:
        filtered.append(row)

print(f"\nTotal: {len(filtered)} records\n")

# Print each record with key info
for i, r in enumerate(filtered, 1):
    post_code = r[col_map.get('岗位编号', 3)]
    post_name = r[col_map.get('岗位名称', 4)]
    do_job = r[col_map.get('从事工作', 5)]
    unit = r[col_map.get('大单位', 0)]
    need_unit = r[col_map.get('用人单位', 1)]
    subject = r[col_map.get('学科专业', 8)]
    location = r[col_map.get('工作地点', 13)]
    print(f"{i:3d}. [{post_code}] {post_name} | {do_job} | {unit}/{need_unit} | {location}")
    print(f"     专业: {subject}")

# Stats
print("\n=== 工作地点分布 ===")
locations = Counter()
for r in filtered:
    loc = r[col_map.get('工作地点', 13)]
    if loc:
        locations[str(loc)] += 1
for loc, n in locations.most_common(15):
    print(f"  {loc}: {n}条")

# Save
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "计算机科学与技术"

header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin", color="D9DEE7"),
    right=Side(style="thin", color="D9DEE7"),
    top=Side(style="thin", color="D9DEE7"),
    bottom=Side(style="thin", color="D9DEE7"),
)

for col_idx, name in enumerate(headers, 1):
    if name:
        cell = ws2.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

for row_idx, record in enumerate(filtered, 2):
    for col_idx, val in enumerate(record, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = cell_align
        cell.border = thin_border

col_widths = [14,22,14,14,18,14,8,16,40,28,30,10,10,14,22]
for i, w in enumerate(col_widths, 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = ws2.dimensions

output = 'd:\\PycharmProjects\\CanvasAnvil\\筛选_计算机科学与技术_only.xlsx'
wb2.save(output)
print(f"\n文件已保存: {output}")
