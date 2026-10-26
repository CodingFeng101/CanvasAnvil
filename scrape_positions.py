import requests
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import time
import sys

BASE_URL = "http://211.166.6.109:9998/enroll/post/listVisitor"
PAGE_SIZE = 100
TOTAL = 5194

COLUMNS = [
    ("大单位", "unit"),
    ("用人单位", "needHandsUnit"),
    ("岗位类别", "jobProperties"),
    ("岗位编号", "postCode"),
    ("岗位名称", "post"),
    ("从事工作", "doJob"),
    ("性别", "isSex"),
    ("学历层次", "educationBackground"),
    ("学科专业", "subject"),
    ("院校及学科层次", "isDoubletop"),
    ("其他要求", "otherReq"),
    ("招录数量", "quantity"),
    ("已报名人数", "postNum"),
    ("工作地点", "forceStation"),
    ("咨询电话", "telPhone"),
]

GENDER_MAP = {"1": "男", "2": "女", "0": "不限"}
DOUBLE_TOP_MAP = {"1": "“双一流”建设高校及建设学科", "0": ""}

def fetch_page(page):
    params = {"page": page, "limit": PAGE_SIZE, "queryStr": ""}
    resp = requests.get(BASE_URL, params=params, timeout=30)
    resp.raise_for_status()
    return resp.json()

def main():
    all_records = []
    total_pages = (TOTAL + PAGE_SIZE - 1) // PAGE_SIZE

    print(f"开始爬取，共 {total_pages} 页，每页 {PAGE_SIZE} 条")

    for page in range(1, total_pages + 1):
        try:
            result = fetch_page(page)
            if result.get("code") == 200:
                records = result["data"]["data"]
                all_records.extend(records)
                actual_total = result["data"]["total"]
                print(f"  第 {page}/{total_pages} 页: 获取 {len(records)} 条，累计 {len(all_records)}/{actual_total}")
            else:
                print(f"  第 {page} 页返回异常: {result.get('msg')}")
        except Exception as e:
            print(f"  第 {page} 页出错: {e}")
            time.sleep(2)
            try:
                result = fetch_page(page)
                if result.get("code") == 200:
                    records = result["data"]["data"]
                    all_records.extend(records)
                    print(f"  第 {page} 页重试成功: 获取 {len(records)} 条")
            except Exception as e2:
                print(f"  第 {page} 页重试失败: {e2}")

        if page < total_pages:
            time.sleep(0.3)

    print(f"\n爬取完成，共获取 {len(all_records)} 条记录")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "招录岗位"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, (display_name, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=display_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, record in enumerate(all_records, 2):
        for col_idx, (_, field) in enumerate(COLUMNS, 1):
            value = record.get(field, "")
            if field == "isSex":
                value = GENDER_MAP.get(str(value), str(value))
            elif field == "isDoubletop":
                value = DOUBLE_TOP_MAP.get(str(value), str(value))
            elif field == "quantity":
                value = record.get("quantity", 0)
            elif field == "otherReq":
                value = record.get("otherReq", "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_align
            cell.border = thin_border

    col_widths = [14, 22, 14, 14, 18, 14, 8, 16, 40, 28, 30, 10, 10, 14, 22]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output_path = "d:\\PycharmProjects\\CanvasAnvil\\招录岗位数据.xlsx"
    wb.save(output_path)
    print(f"\nExcel 文件已保存至: {output_path}")
    print(f"数据行数: {len(all_records)}")

if __name__ == "__main__":
    main()
