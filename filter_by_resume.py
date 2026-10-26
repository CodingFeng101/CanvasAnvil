import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import Counter

wb = openpyxl.load_workbook('d:\\PycharmProjects\\CanvasAnvil\\招录岗位数据.xlsx')
ws = wb.active

computer_majors = {
    '计算机科学与技术', '计算机技术', '计算机系统', '计算机',
    '软件工程', '软件技术',
    '网络工程', '网络空间安全', '网络与信息安全',
    '网路工程', '网路空间安全', '网路与信息安全',
    '信息安全', '保密技术', '保密管理',
    '密码科学与技术', '密码',
    '数据科学与大数据技术', '大数据技术与工程', '大数据技术',
    '人工智能',
    '智能科学与技术', '智能无人系统技术',
    '电子信息工程', '电子信息科学与技术',
    '电子科学与技术',
    '通信工程',
    '信息工程',
    '信息对抗技术',
    '微电子科学与工程', '微电子学与固体电子学',
    '光电信息科学与工程',
    '电子与计算机工程',
    '空间信息与数字技术',
    '电磁场与无线技术', '电波传播与天线',
    '数字媒体技术', '新媒体技术',
    '机器人工程',
    '信息与计算科学',
}

def is_computer_related(subject):
    if not subject:
        return False
    s = str(subject)
    parts = s.replace('，', '、').replace(',', '、').split('、')
    for part in parts:
        part = part.strip()
        if part in computer_majors:
            return True
        for major in computer_majors:
            if part.startswith(major + '（') or part.startswith(major + '('):
                return True
    return False

def has_exact_cs(subject):
    if not subject:
        return False
    parts = str(subject).replace('，', '、').replace(',', '、').split('、')
    for part in parts:
        part = part.strip()
        if part == '计算机科学与技术' or part.startswith('计算机科学与技术（') or part.startswith('计算机科学与技术('):
            return True
    return False

# Filter all 182 records
filtered_all = []
filtered_cs = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[6] == '男' and row[7] == '本科生' and not row[9]:
        if is_computer_related(row[8]):
            filtered_all.append(row)
            if has_exact_cs(row[8]):
                filtered_cs.append(row)

print(f"计算机相关专业总数: {len(filtered_all)} 条")
print(f"其中明确包含'计算机科学与技术'的: {len(filtered_cs)} 条")

# Show all unique subject combinations that include CS
print()
print("=== 包含'计算机科学与技术'的学科专业组合 ===")
subjects_cs = Counter()
for r in filtered_cs:
    subjects_cs[r[8]] += 1
for s, n in subjects_cs.most_common():
    print(f"  [{n:3d}条] {s}")

# The remaining records that DON'T include CS specifically
print()
print("=== 不含'计算机科学与技术'的计算机相关专业组合 ===")
subjects_non_cs = Counter()
for r in filtered_all:
    if not has_exact_cs(r[8]):
        subjects_non_cs[r[8]] += 1
for s, n in subjects_non_cs.most_common():
    print(f"  [{n:3d}条] {s}")

# Save the CS-specific results
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "计算机科学与技术岗位"

header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

headers = ['大单位','用人单位','岗位类别','岗位编号','岗位名称','从事工作','性别',
           '学历层次','学科专业','院校及学科层次','招录数量','已报名人数','工作地点','咨询电话']
for col_idx, name in enumerate(headers, 1):
    cell = ws2.cell(row=1, column=col_idx, value=name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

for row_idx, record in enumerate(filtered_cs, 2):
    out_cols = [0,1,2,3,4,5,6,7,8,9,11,12,13,14]
    for col_idx, src_idx in enumerate(out_cols, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=record[src_idx])
        cell.alignment = cell_align
        cell.border = thin_border

col_widths = [14,22,14,14,18,14,8,16,40,28,10,10,14,22]
for i, w in enumerate(col_widths, 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = ws2.dimensions

output = 'd:\\PycharmProjects\\CanvasAnvil\\筛选_计算机科学与技术_可报岗位.xlsx'
wb2.save(output)
print(f"\nExcel 文件已保存至: {output}")
print(f"数据行数: {len(filtered_cs)}")

# Work location stats
locations = Counter()
for r in filtered_cs:
    locations[r[13]] += 1
print(f"\n工作地点分布:")
for loc, n in locations.most_common(15):
    print(f"  {loc}: {n}条")
