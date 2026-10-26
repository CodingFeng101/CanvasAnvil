import requests
import time
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "http://211.166.6.109:9998/enroll/post/listVisitor"
PAGE_SIZE = 100

session = requests.Session()
session.trust_env = False

def fetch_all():
    all_records = []
    total_pages = None
    for page in range(1, 60):
        params = {"page": page, "limit": PAGE_SIZE, "queryStr": ""}
        try:
            resp = session.get(BASE_URL, params=params, timeout=30)
            data = resp.json()
            if data.get("code") == 200:
                records = data["data"]["data"]
                if total_pages is None:
                    total = data["data"]["total"]
                    total_pages = (total + PAGE_SIZE - 1) // PAGE_SIZE
                    print(f"总记录数: {total}, 总页数: {total_pages}")
                all_records.extend(records)
                print(f"已爬取第 {page}/{total_pages} 页, 累计 {len(all_records)} 条", end="\r")
                if page >= total_pages:
                    break
            else:
                print(f"第{page}页返回错误: {data}")
                break
        except Exception as e:
            print(f"第{page}页异常: {e}")
            time.sleep(1)
    return all_records

def save_excel(records, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "招录岗位"

    headers = ["大单位", "用人单位", "岗位类别", "岗位编号", "岗位名称", "从事工作",
               "性别", "学历层次", "学科专业", "院校及学科层次", "其他要求",
               "招录数量", "已报名人数", "工作地点", "咨询电话"]

    # Header style
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Data rows
    data_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for i, r in enumerate(records, 2):
        # Map fields
        is_double = str(r.get("isDoubletop", ""))
        if is_double == "1":
            edu_level = "双一流"
        elif is_double == "0":
            edu_level = "一本以上"
        else:
            edu_level = is_double

        sr = str(r.get("specialRequirement", ""))
        if sr == "1":
            other_req = "应为\"双一流\"建设高校毕业生"
        else:
            other_req = "—"

        sex = str(r.get("isSex", ""))
        if sex == "1":
            sex_str = "男"
        elif sex == "0":
            sex_str = "不限"
        elif sex == "2":
            sex_str = "女"
        else:
            sex_str = sex

        row_data = [
            r.get("unit", ""),
            r.get("needHandsUnit", ""),
            r.get("jobProperties", ""),
            r.get("postCode", ""),
            r.get("post", ""),
            r.get("doJob", ""),
            sex_str,
            r.get("educationBackground", ""),
            r.get("subject", ""),
            edu_level,
            other_req,
            r.get("quantity", 0),
            r.get("postNum", 0),
            r.get("forceStation", ""),
            r.get("telephone", ""),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.alignment = data_align
            cell.border = thin_border

    # Column widths
    widths = [10, 18, 14, 14, 12, 12, 6, 12, 30, 12, 18, 8, 10, 14, 22]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # Freeze header
    ws.freeze_panes = "A2"
    # Auto filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records)+1}"

    wb.save(filename)
    print(f"\n保存到 {filename}, 共 {len(records)} 条")

def main():
    print("开始爬取全部岗位数据...")
    records = fetch_all()
    print(f"\n爬取完成! 共 {len(records)} 条")

    # Save JSON
    with open("all_positions.json", "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    print("JSON已保存: all_positions.json")

    # Save Excel
    save_excel(records, "招录岗位数据_最新.xlsx")
    print("Excel已保存: 招录岗位数据_最新.xlsx")

if __name__ == "__main__":
    main()
