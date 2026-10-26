import openpyxl

wb = openpyxl.load_workbook('d:\\PycharmProjects\\CanvasAnvil\\招录岗位数据.xlsx')
ws = wb.active

col_names = ['大单位','用人单位','岗位类别','岗位编号','岗位名称','从事工作','性别',
             '学历层次','学科专业','院校及学科层次','其他要求','招录数量','已报名人数','工作地点','咨询电话']

# Check all columns for batch-related keywords
found = []
for row in ws.iter_rows(min_row=2, values_only=True):
    for col_idx, val in enumerate(row):
        if val and any(kw in str(val) for kw in ['一批', '二批', '一本', '二本', '批次']):
            found.append((row[3], row[4], col_idx, val))

print(f'包含"一批/二批/一本/二本/批次"的记录数: {len(found)}')
if found:
    for r in found[:30]:
        print(f'  岗位编号:{r[0]} | 岗位名称:{r[1]} | {col_names[r[2]]}: {r[3]}')
else:
    print('  无相关记录')

# Also check unique values in 院校及学科层次 column (index 9)
print()
print('=== 院校及学科层次 所有唯一值 ===')
levels = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[9]:
        levels.add(str(row[9]))

for v in sorted(levels):
    print(f'  {v}')

# Also check 其他要求 column (index 10) for any education batch requirements
print()
print('=== 其他要求 所有唯一值 ===')
others = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[10]:
        others.add(str(row[10]))

for v in sorted(others):
    print(f'  {v}')
