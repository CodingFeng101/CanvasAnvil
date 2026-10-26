import openpyxl
from collections import Counter

wb = openpyxl.load_workbook('d:\\PycharmProjects\\CanvasAnvil\\招录岗位数据.xlsx')
ws = wb.active

computer_majors = {
    '计算机科学与技术', '计算机技术', '计算机系统', '计算机',
    '软件工程', '软件技术',
    '网络工程', '网络空间安全', '网络与信息安全',
    '网路工程', '网路空间安全', '网路与信息安全',
    '信息安全', '保密技术', '保密管理',
    '密码科学与技术', '密码',
    '数据科学与大数据技术', '大数据技术与工程', '大数据技术',
    '人工智能',
    '智能科学与技术', '智能无人系统技术',
    '电子信息工程', '电子信息科学与技术',
    '电子科学与技术',
    '通信工程',
    '信息工程',
    '信息对抗技术',
    '微电子科学与工程', '微电子学与固体电子学',
    '光电信息科学与工程',
    '电子与计算机工程',
    '空间信息与数字技术',
    '电磁场与无线技术', '电波传播与天线',
    '数字媒体技术', '新媒体技术',
    '机器人工程',
    '信息与计算科学',
}

def is_matched(subject):
    if not subject:
        return False
    s = str(subject)
    parts = s.replace('，', '、').replace(',', '、').split('、')
    for part in parts:
        part = part.strip()
        if part in computer_majors:
            return True
        for major in computer_majors:
            if part.startswith(major + '（') or part.startswith(major + '('):
                return True
    return False

# Get the 1425 filtered records (男+本科生+非双一流)
filtered_subjects = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[6] == '男' and row[7] == '本科生' and not row[9]:
        filtered_subjects[row[8]] += 1

# Broad keywords to search for potentially missed majors
broad_keywords = [
    '计算机', '软件', '网络', '网路', '信息安全', '保密', '密码',
    '大数据', '数据科学', '人工智能', '智能科学', '智能无人',
    '电子信息', '电子科学', '通信', '信息工程', '信息对抗',
    '微电子', '光电信息', '电磁场', '电波传播', '机器人',
    '电子与计算机', '空间信息与数字', '数字媒体', '新媒体技术',
    '信息与计算科学', '无人驾驶航空器',
    '集成电路', '物联网', '半导体', '芯片',
    '信息管理', '信息系统', '遥感', '地理空间',
    '控制', '自动化', '电气', '智能装备', '智能车辆',
    '智能建造', '智能飞行', '系统科学', '管理科学',
    '网路', '网络与新媒体', '网络空间',
]

# Find subjects that contain a keyword but were NOT matched
print('=== 包含关键词但未被匹配的学科专业（可能遗漏）===')
found_suspects = False
seen = set()
for s, n in filtered_subjects.most_common():
    if not is_matched(s) and s not in seen:
        seen.add(s)
        parts = str(s).replace('，', '、').replace(',', '、').split('、')
        for part in parts:
            part = part.strip()
            matched_kw = None
            for kw in broad_keywords:
                if kw in part:
                    matched_kw = kw
                    break
            if matched_kw:
                print(f'  [{n:3d}条] 关键词"{matched_kw}" -> "{part}"')
                print(f'         完整: {s}')
                found_suspects = True

if not found_suspects:
    print('  无遗漏')

# Also verify the count
matched_count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[6] == '男' and row[7] == '本科生' and not row[9]:
        if is_matched(row[8]):
            matched_count += 1

print(f'\n验证: 匹配 {matched_count} 条 (之前结果 182 条)')
